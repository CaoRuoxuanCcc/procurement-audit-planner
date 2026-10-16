"""Shared helpers for the QA checks in scripts/checks/.

Each check prints one `FAIL <item>` line per failure, then one final line
`RESULT {json}` with check, status (PASS / FAIL / STOP), checked and failures;
run_checks.py reads that line. Exit: 0 PASS, 1 FAIL, 2 STOP (an input is
missing or unreadable).

The checks re-read the rendered documents and the raw sources themselves; they
do not import the builders, so a builder fault cannot hide itself.
"""
import argparse
import csv
import json
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SHEET = "Data to publish"
SPEND_COLUMNS = ["Service", "Expenditure category", "Narrative", "Date", "Net amount", "Supplier name"]
PACK_DOCS = ("planning-memo.md", "risk-control-matrix.md", "audit-program.md")
ROW_ID_RE = re.compile(r"[^\s|`\"',;]+\.xlsx\|" + re.escape(SHEET) + r"\|R\d+")


class Stop(Exception):
    """Missing or malformed input."""


def args(description):
    ap = argparse.ArgumentParser(description=description)
    ap.add_argument("--out-dir", default=os.path.join(ROOT, "outputs"))
    ap.add_argument("--data-dir", default=os.path.join(ROOT, "data"))
    return ap.parse_args()


def need(path):
    if not os.path.exists(path):
        raise Stop(f"{os.path.relpath(path, ROOT) if path.startswith(ROOT) else path} not found")
    return path


def read_text(path):
    with open(need(path), encoding="utf-8") as f:
        return f.read()


def read_json(path):
    try:
        return json.loads(read_text(path))
    except json.JSONDecodeError as e:
        raise Stop(f"{path} is not valid JSON: {e}")


def read_csv(path):
    with open(need(path), newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def md_rows(text, section=None):
    """Cells of every markdown table row (header and rule rows dropped), optionally under one '## ' section."""
    if section is not None:
        m = re.search(rf"^## {re.escape(section)}\s*$(.*?)(?=^## |\Z)", text, re.S | re.M)
        if not m:
            return None
        text = m.group(1)
    rows, header_seen = [], False
    for line in text.splitlines():
        if not line.startswith("|"):
            header_seen = False
            continue
        cells = split_row(line)
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells if c):
            header_seen = True
            continue
        if not header_seen:
            continue  # the header row comes before the rule row
        rows.append(cells)
    return rows


def split_row(line):
    s = line.strip()
    s = s[1:] if s.startswith("|") else s
    s = s[:-1] if s.endswith("|") and not s.endswith("\\|") else s
    cells, cur, i = [], "", 0
    while i < len(s):
        if s[i] == "\\" and i + 1 < len(s) and s[i + 1] == "|":
            cur += "|"
            i += 2
            continue
        if s[i] == "|":
            cells.append(cur.strip())
            cur = ""
        else:
            cur += s[i]
        i += 1
    cells.append(cur.strip())
    return cells


def cents(x):
    return round(float(x) * 100)


def load_workbooks(data_dir):
    """{file: {excel_row: (supplier, amount)}} read straight from data/spend/*.xlsx."""
    import openpyxl
    spend_dir = need(os.path.join(data_dir, "spend"))
    out = {}
    for fn in sorted(os.listdir(spend_dir)):
        if not fn.endswith(".xlsx"):
            continue
        wb = openpyxl.load_workbook(os.path.join(spend_dir, fn), read_only=True, data_only=True)
        if SHEET not in wb.sheetnames:
            raise Stop(f"{fn}: sheet {SHEET!r} not found")
        ws = wb[SHEET]
        pos = None
        for row in ws.iter_rows(min_row=1, max_row=20):
            vals = ["" if c.value is None else str(c.value).strip() for c in row]
            if all(c in vals for c in SPEND_COLUMNS):
                hdr, pos = row[0].row, {c: vals.index(c) for c in SPEND_COLUMNS}
                break
        if pos is None:
            raise Stop(f"{fn}: no header row with the six spend columns")
        rows = {}
        for row in ws.iter_rows(min_row=hdr + 1):
            cells = {c.row for c in row if getattr(c, "row", None)}
            if not cells:
                continue
            r = cells.pop()
            get = lambda c: row[pos[c]].value if pos[c] < len(row) else None  # noqa: E731
            sup, amt = get("Supplier name"), get("Net amount")
            rows[r] = ("" if sup is None else str(sup).strip(), amt)
        out[fn] = rows
        wb.close()
    return out


def report(check, checked, failures, note=""):
    for f in failures:
        print(f"FAIL {f}")
    status = "PASS" if not failures else "FAIL"
    print(f"{check}: {status}, {checked} checked, {len(failures)} failure(s)" + (f"; {note}" if note else ""))
    print("RESULT " + json.dumps({"check": check, "status": status, "checked": checked,
                                  "failures": failures, "note": note}, ensure_ascii=False))
    return 0 if not failures else 1


def run(check, fn):
    """Call fn() -> (checked, failures, note); turn Stop into a STOP result."""
    try:
        checked, failures, note = fn()
    except Stop as e:
        print(f"STOP {e}")
        print("RESULT " + json.dumps({"check": check, "status": "STOP", "checked": 0,
                                      "failures": [f"STOP: {e}"], "note": ""}, ensure_ascii=False))
        return 2
    return report(check, checked, failures, note)


if __name__ == "__main__":
    sys.exit("qa_common.py is a helper module; run run_checks.py")
