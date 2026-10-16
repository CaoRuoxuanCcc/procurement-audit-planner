"""Clean the monthly spend workbooks.

Reads: data/spend/*.xlsx (sheet 'Data to publish', via openpyxl), outputs/rules.json
       (validated only), scripts/baseline_values.json (Service/Narrative values
       seen in the 2026 P02-P04 extract; created from those files if absent).
Writes: outputs/analytics-full/spend_clean.csv, exclusions.csv, tag_mapping.csv,
        clean_summary.json.
"""
import csv
import glob
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl  # noqa: E402

from common import (FULL, SHEET, SPEND_COLUMNS, SPEND_DIR, StopError,  # noqa: E402
                    load_rules, load_tiers, normalise_supplier, sha256, write_json)

BASELINE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        "baseline_values.json")

PENSION_NARR = {"LGPS Employers Control", "LGPS Employees Control",
                "Teachers LGPS ERS Control", "Employers LGPS Deficit Contribution",
                "Prudential AVC Teachers", "Court Order Control",
                "Council Tax Court Orders", "Rent Rebates"}
PENSION_SUBSTR = re.compile(r"LGPS|PENSION|AVC", re.IGNORECASE)
PUBLIC_NARR = {"Joint Arrangements (with Other Local Authorities or NHS)",
               "Payments to NHS (excl Joint Arrangement)",
               "Payments to Other Local Authorities (excl Joint Arrangements)"}
PUBLIC_NAME_RE = re.compile(
    r"\b(BOROUGH|COUNTY|DISTRICT|CITY|PARISH|TOWN) COUNCIL\b|\bCOUNCIL$|\bNHS\b|"
    r"\bICB\b|INTEGRATED CARE BOARD|\bHMRC\b|REVENUE AND CUSTOMS|\bPOLICE\b|"
    r"FIRE AND RESCUE|ENVIRONMENT AGENCY|^DEPARTMENT FOR\b|^MINISTRY OF\b|"
    r"^ROYAL BOROUGH OF\b")
CARE_SERVICES = {"Adult Social Care", "Adult Social Care - Transformation Projects",
                 "Children's Social Care", "Education & SEND",
                 "Education (DSG Funded) & Other Education Grants"}
CARE_NARR = {"Private Contractors", "Voluntary Associations", "Other agencies",
             "Allowances for provision care", "Adult Placement Allowances",
             "Payments to Carers", "Direct Payment", "ASC Deferred Debt",
             "Private Contractors Additional Cost"}
TAG_MAP_NARR = {"Grants", "Agency & Temporary Staff"} | CARE_NARR


def s(v):
    return "" if v is None else str(v).strip()


def find_header(ws, fname):
    for row in ws.iter_rows(min_row=1, max_row=20):
        vals = [s(c.value) for c in row]
        if all(col in vals for col in SPEND_COLUMNS):
            return row[0].row, {col: vals.index(col) for col in SPEND_COLUMNS}
    raise StopError(f"{fname}: no header row with all six columns in first 20 rows")


def read_file(path):
    fname = os.path.basename(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise StopError(f"{fname}: sheet '{SHEET}' not found")
    ws = wb[SHEET]
    hdr_row, pos = find_header(ws, fname)
    rows = []
    for row in ws.iter_rows(min_row=hdr_row + 1):
        if all(c.value is None or s(c.value) == "" for c in row):
            continue
        excel_row = next(c.row for c in row if hasattr(c, "row") and c.row)
        rid = f"{fname}|{SHEET}|R{excel_row}"
        get = lambda col: row[pos[col]].value if pos[col] < len(row) else None  # noqa: E731
        dv = get("Date")
        if isinstance(dv, datetime):
            ts = dv
        elif isinstance(dv, str):
            try:
                ts = datetime.strptime(dv.strip(), "%m/%d/%y")
            except ValueError:
                raise StopError(f"{rid}: unparseable text date")
        else:
            raise StopError(f"{rid}: Date cell of unsupported type {type(dv).__name__}")
        amt = get("Net amount")
        if isinstance(amt, bool) or not isinstance(amt, (int, float)):
            raise StopError(f"{rid}: Net amount is not numeric ({type(amt).__name__})")
        rows.append({"row_id": rid, "file": fname, "excel_row": excel_row,
                     "service": s(get("Service")), "category": s(get("Expenditure category")),
                     "narrative": s(get("Narrative")), "timestamp": ts,
                     "date": ts.date(), "amount": float(amt),
                     "supplier": s(get("Supplier name"))})
    wb.close()
    return hdr_row, rows


def exclusion(r):
    sup = r["supplier"]
    if sup.lower() == "redacted":
        return "redacted_supplier", ""
    n = r["narrative"]
    if n in PENSION_NARR or PENSION_SUBSTR.search(n):
        return "pension_or_statutory", ""
    if r["category"] == "" and n not in TAG_MAP_NARR:
        return "pension_or_statutory", "blank category"
    if n in PUBLIC_NARR:
        return "public_body", "narrative"
    if PUBLIC_NAME_RE.search(sup.upper()):
        return "public_body", "name pattern"
    if r["amount"] <= 0:
        return "non_positive_amount", ""
    return None, ""


def tag(r):
    n = r["narrative"]
    if n == "Grants":
        return "grant"
    if n == "Agency & Temporary Staff":
        return "agency_staff"
    if r["service"] in CARE_SERVICES and n in CARE_NARR:
        return "care_or_education_placement"
    if r["category"] == "Premises":
        return "premises"
    return "general_procurement"


def main():
    rules = load_rules()
    load_tiers(rules)  # validates the threshold clauses exist
    files = sorted(glob.glob(os.path.join(SPEND_DIR, "*.xlsx")))
    if not files:
        raise StopError("no data/spend/*.xlsx files")
    per_file, all_rows = [], []
    for f in files:
        hdr, rows = read_file(f)
        if not rows:
            raise StopError(f"{os.path.basename(f)}: no data rows")
        per_file.append({"file": os.path.basename(f), "sha256": sha256(f),
                         "header_excel_row": hdr, "rows": len(rows),
                         "date_min": str(min(r["date"] for r in rows)),
                         "date_max": str(max(r["date"] for r in rows))})
        all_rows.extend(rows)

    # cross-file duplicates (counted, not dropped)
    key_files = defaultdict(set)
    for r in all_rows:
        key_files[(r["timestamp"], r["amount"], r["supplier"], r["narrative"])].add(r["file"])
    cross = sum(1 for r in all_rows
                if len(key_files[(r["timestamp"], r["amount"], r["supplier"], r["narrative"])]) > 1)

    os.makedirs(FULL, exist_ok=True)
    kept, excl = [], []
    for r in all_rows:
        reason, note = exclusion(r)
        if reason:
            r["reason"], r["note"] = reason, note
            excl.append(r)
        else:
            r["tag"] = tag(r)
            r["supplier_norm"] = normalise_supplier(r["supplier"])
            kept.append(r)

    with open(os.path.join(FULL, "exclusions.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["row_id", "file", "reason", "note", "amount", "narrative",
                    "category", "supplier"])
        for r in excl:
            w.writerow([r["row_id"], r["file"], r["reason"], r["note"], repr(r["amount"]),
                        r["narrative"], r["category"], r["supplier"]])

    tm = defaultdict(lambda: [0, 0.0])
    for r in kept:
        k = (r["service"], r["category"], r["narrative"], r["tag"])
        tm[k][0] += 1
        tm[k][1] += r["amount"]
    with open(os.path.join(FULL, "tag_mapping.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["service", "category", "narrative", "tag", "rows", "total_gbp"])
        for k in sorted(tm):
            w.writerow(list(k) + [tm[k][0], f"{tm[k][1]:.2f}"])

    cols = ["row_id", "file", "excel_row", "service", "category", "narrative",
            "timestamp", "date", "amount", "supplier", "supplier_norm", "tag"]
    with open(os.path.join(FULL, "spend_clean.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for r in kept:
            w.writerow([r["row_id"], r["file"], r["excel_row"], r["service"], r["category"],
                        r["narrative"], r["timestamp"].isoformat(), r["date"].isoformat(),
                        repr(r["amount"]), r["supplier"], r["supplier_norm"], r["tag"]])

    # baseline of Service / Narrative values from the 2026 P02-P04 extract
    import json
    if not os.path.exists(BASELINE):
        base_rows = [r for r in all_rows if re.search(r"2026_P0[234]", r["file"])]
        with open(BASELINE, "w") as f:
            json.dump({"source": "2026 P02-P04 extract",
                       "service": sorted({r["service"] for r in base_rows}),
                       "narrative": sorted({r["narrative"] for r in base_rows})}, f, indent=1)
    with open(BASELINE) as f:
        base = json.load(f)
    unseen = {"service": sorted({r["service"] for r in all_rows} - set(base["service"])),
              "narrative": sorted({r["narrative"] for r in all_rows} - set(base["narrative"]))}

    excluded = defaultdict(lambda: {"count": 0, "total_gbp": 0.0})
    for r in excl:
        excluded[r["reason"]]["count"] += 1
        excluded[r["reason"]]["total_gbp"] += r["amount"]
    by_tag = defaultdict(lambda: {"count": 0, "total_gbp": 0.0})
    for r in kept:
        by_tag[r["tag"]]["count"] += 1
        by_tag[r["tag"]]["total_gbp"] += r["amount"]
    for pf in per_file:
        fk = [r for r in kept if r["file"] == pf["file"]]
        fe = [r for r in excl if r["file"] == pf["file"]]
        pf["kept_rows"], pf["excluded_rows"] = len(fk), len(fe)
        pf["kept_gbp"] = round(sum(r["amount"] for r in fk), 2)
        pf["excluded_gbp"] = round(sum(r["amount"] for r in fe), 2)
        pf["raw_gbp"] = round(sum(r["amount"] for r in all_rows if r["file"] == pf["file"]), 2)
        pf["excluded_by_reason"] = dict(Counter(r["reason"] for r in fe))

    public_names = Counter(r["supplier"] for r in excl if r["reason"] == "public_body")
    summary = {
        "spend_files": per_file,
        "window": {"start": str(min(r["date"] for r in kept)),
                   "end": str(max(r["date"] for r in kept)),
                   "months": len(files),
                   "annualisation_factor": 12 / len(files)},
        "rows_read": len(all_rows),
        "distinct_supplier_strings": len({r["supplier"] for r in all_rows}),
        "excluded": {k: {"count": v["count"], "total_gbp": round(v["total_gbp"], 2)}
                     for k, v in sorted(excluded.items())},
        "blank_category_rows_read": sum(1 for r in all_rows if r["category"] == ""),
        "blank_category_rows_kept": sum(1 for r in kept if r["category"] == ""),
        "rows_tested": len(kept),
        "total_tested_gbp": round(sum(r["amount"] for r in kept), 2),
        "by_tag": {k: {"count": v["count"], "total_gbp": round(v["total_gbp"], 2)}
                   for k, v in sorted(by_tag.items())},
        "cross_file_duplicate_rows": cross,
        "same_timestamp_dup_rows_within_file": sum(
            c for c in Counter((r["file"], r["timestamp"], r["amount"], r["supplier"],
                                r["narrative"]) for r in all_rows).values() if c > 1),
        "public_body_distinct_names_caught": len(public_names),
        "unseen_values": unseen,
        "unseen_values_baseline": "scripts/baseline_values.json (2026 P02-P04 extract)",
        "negative_or_zero_rows": sum(1 for r in all_rows if r["amount"] <= 0),
        "min_kept_amount": min(r["amount"] for r in kept),
    }
    write_json(os.path.join(FULL, "clean_summary.json"), summary)
    print(f"files={len(files)} rows_read={len(all_rows)} kept={len(kept)} "
          f"excluded={len(excl)} window={summary['window']}")
    for pf in per_file:
        print(f"  {pf['file']}: header R{pf['header_excel_row']} rows={pf['rows']} "
              f"kept={pf['kept_rows']} excl={pf['excluded_rows']} {pf['date_min']}..{pf['date_max']}")
    print("excluded:", summary["excluded"])
    print("by_tag:", summary["by_tag"])
    print("cross_file_dup_rows:", cross, "unseen:", unseen)


if __name__ == "__main__":
    try:
        main()
    except StopError as e:
        print(f"STOP: {e}", file=sys.stderr)
        sys.exit(2)
