"""Independent self-check of the spend-analyst outputs.

Reads: the RAW data/spend/*.xlsx (openpyxl), the RAW data/contracts.csv (csv
       module), outputs/rules.json, outputs/analytics.json and every CSV in
       outputs/analytics-full/; scans scripts/*.py.
Writes: outputs/analytics-full/self_check.json, and updates only the
        'self_check' field of outputs/analytics.json. Exits 1 on any failure.
Imports from common only normalise_supplier, trading_parts and BANNED_WORDS.
"""
import csv
import glob
import json
import os
import re
import sys
from datetime import datetime
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl  # noqa: E402

from common import BANNED_WORDS, normalise_supplier, trading_parts  # noqa: E402,F401

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FULL = os.path.join(ROOT, "outputs", "analytics-full")
AJ = os.path.join(ROOT, "outputs", "analytics.json")
SHEET = "Data to publish"
COLS = ["Service", "Expenditure category", "Narrative", "Date", "Net amount", "Supplier name"]
SUP_COL = "Supplier [Name|Address|Ref type|Ref Number|Is SME|Is VCSE]"
results = []


def record(name, ok, detail):
    results.append({"check": name, "result": "PASS" if ok else "FAIL", "detail": detail})


def cents(x):
    return int(round(float(x) * 100))


# ---------------------------------------------------------------- raw spend
def load_raw():
    raw = {}
    for path in sorted(glob.glob(os.path.join(ROOT, "data", "spend", "*.xlsx"))):
        fn = os.path.basename(path)
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True)[SHEET]
        hdr, pos = None, None
        for row in ws.iter_rows(min_row=1, max_row=20):
            vals = ["" if c.value is None else str(c.value).strip() for c in row]
            if all(c in vals for c in COLS):
                hdr, pos = row[0].row, {c: vals.index(c) for c in COLS}
                break
        rows = {}
        for row in ws.iter_rows(min_row=hdr + 1):
            if all(c.value is None or str(c.value).strip() == "" for c in row):
                continue
            rnum = next(c.row for c in row if getattr(c, "row", None))
            d = row[pos["Date"]].value
            if isinstance(d, str):
                d = datetime.strptime(d.strip(), "%m/%d/%y")
            sup = "" if row[pos["Supplier name"]].value is None else str(row[pos["Supplier name"]].value).strip()
            rows[rnum] = {"amount": row[pos["Net amount"]].value, "date": d.date(), "ts": d,
                          "supplier": sup, "norm": normalise_supplier(sup)}
        raw[fn] = rows
    return raw


def resolve(raw, rid):
    parts = rid.split("|")
    if len(parts) != 3 or parts[1] != SHEET or not parts[2].startswith("R"):
        return None
    return raw.get(parts[0], {}).get(int(parts[2][1:]))


def read_csv(name):
    with open(os.path.join(FULL, name), newline="") as f:
        return list(csv.DictReader(f))


# ---------------------------------------------------------------- thresholds
def thresholds_from_rules():
    with open(os.path.join(ROOT, "outputs", "rules.json")) as f:
        rules = json.load(f)
    by = {r["clause"]: r for r in rules["rules"]}
    out = {}
    for cl in ["App B row A", "App B row B1", "App B row B1 (advertising)", "App B row C"]:
        r = by[cl]
        if r.get("threshold_low") is not None:
            out[r["threshold_low"]] = bool(r.get("threshold_low_inclusive"))
    return out  # {T: inclusive}


def enters(a, T, inc):
    return a >= T if inc else a > T


# ---------------------------------------------------------------- raw contracts
def council_notices():
    path = os.path.join(ROOT, "data", "contracts.csv")
    if not os.path.exists(path):
        return None
    out = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            if r["Organisation Name"].strip().lower() not in ("west berkshire council", "west berkshire"):
                continue
            raw = r[SUP_COL].strip()
            raw = raw[1:] if raw.startswith("[") else raw
            raw = raw[:-1] if raw.endswith("]") else raw
            blocks = raw.split("][") if raw else []
            for b in blocks:
                out.append({"notice_id": r["Notice Identifier"], "norm": normalise_supplier(b.split("|")[0]),
                            "end": r["Contract end date"].strip(), "start": r["Contract start date"].strip(),
                            "value": r["Awarded Value"].replace(",", "").strip(), "n": len(blocks)})
    return out


def main():
    raw = load_raw()
    with open(AJ) as f:
        aj = json.load(f)
    Ts = thresholds_from_rules()
    nfiles = len(raw)
    factor = 12 / nfiles
    kept = read_csv("spend_clean.csv")
    excl = read_csv("exclusions.csv")

    # 1. reconciliation per file
    for fn, rows in raw.items():
        k = [r for r in kept if r["row_id"].split("|")[0] == fn]
        e = [r for r in excl if r["row_id"].split("|")[0] == fn]
        ids = {r["row_id"] for r in k} | {r["row_id"] for r in e}
        raw_ids = {f"{fn}|{SHEET}|R{n}" for n in rows}
        raw_c = sum(cents(v["amount"]) for v in rows.values())
        kc = sum(cents(resolve(raw, r["row_id"])["amount"]) for r in k)
        ec = sum(cents(resolve(raw, r["row_id"])["amount"]) for r in e)
        kc_out = sum(cents(r["amount"]) for r in k)
        ec_out = sum(cents(r["amount"]) for r in e)
        ok = (len(rows) == len(k) + len(e) and ids == raw_ids and raw_c == kc + ec
              and kc == kc_out and ec == ec_out)
        record(f"reconcile:{fn}", ok, f"raw rows={len(rows)} kept={len(k)} excluded={len(e)}; "
               f"raw GBP={raw_c / 100:.2f} kept={kc_out / 100:.2f} excluded={ec_out / 100:.2f}")
    tot = sum(cents(r["amount"]) for r in kept)
    record("reconcile:analytics_total", tot == cents(aj["cleaning"]["total_tested_gbp"]),
           f"spend_clean total {tot / 100:.2f} vs analytics.json {aj['cleaning']['total_tested_gbp']}")
    excluded_ids = {r["row_id"] for r in excl}
    window_start = min(resolve(raw, r["row_id"])["date"] for r in kept)
    tests = aj["tests"]

    def first(t):
        e = tests.get(t, {})
        if e.get("status") != "ran" or not e.get("top_20"):
            return None
        return e["top_20"][0]

    # 2. re-derive one item per test
    x = first("threshold_clustering")
    if x:
        r = resolve(raw, x["row_id"])
        T = x["T"]
        ok = (r is not None and T in Ts and r["amount"] >= 0.9 * T and not enters(r["amount"], T, Ts[T])
              and cents(r["amount"]) == cents(x["amount"]))
        record("rederive:threshold_clustering", ok, f"{x['row_id']} amount in band below T={T}")

    x = first("split_purchases")
    if x:
        g = [g for g in read_csv("split_purchases.csv") if g["supplier_norm"] == x["supplier_norm"]
             and int(float(g["T"])) == x["T"] and g["first_date"] == x["first_date"]
             and int(g["n_payments"]) == x["n_payments"]][0]
        rows = [resolve(raw, i) for i in g["row_ids"].split(";")]
        T, inc = x["T"], Ts[x["T"]]
        ok = all(r is not None for r in rows)
        norms = {r["norm"] for r in rows}
        ok = ok and norms == {x["supplier_norm"]} and all(not enters(r["amount"], T, inc) for r in rows)
        total = sum(cents(r["amount"]) for r in rows)
        ok = ok and total == cents(x["group_total"]) and enters(total / 100, T, inc)
        # every row must sit in a window of <= 29 days (>= 2 rows) whose sum enters the tier
        dates = sorted(r["date"] for r in rows)
        span = (dates[-1] - dates[0]).days
        covered = set()
        for i, a in enumerate(rows):
            w = [j for j, b in enumerate(rows) if 0 <= (b["date"] - a["date"]).days <= 29]
            if len(w) >= 2 and enters(sum(rows[j]["amount"] for j in w), T, inc):
                covered.update(w)
        ok = ok and len(covered) == len(rows)
        record("rederive:split_purchases", ok,
               f"{x['supplier_norm']} T={T}: {len(rows)} rows, one supplier, all under T, sum "
               f"{total / 100:.2f} = group_total and enters tier; group span {span} days; every "
               f"row lies in a window of at most 29 days whose sum enters the tier")

    notices = council_notices()
    x = first("off_contract")
    if x:
        row = [r for r in read_csv("off_contract.csv") if r["supplier_norm"] == x["supplier_norm"]][0]
        ids = row["row_ids"].split(";")
        rows = [resolve(raw, i) for i in ids]
        all_ids = {f"{fn}|{SHEET}|R{n}" for fn, rr in raw.items() for n, v in rr.items()
                   if v["norm"] == x["supplier_norm"]} - excluded_ids
        tot = sum(r["amount"] for r in rows)
        ok = (all(r and r["norm"] == x["supplier_norm"] for r in rows) and set(ids) == all_ids
              and round(tot * factor, 2) == round(x["annualised_estimate"], 2)
              and notices is not None and x["supplier_norm"] not in {n["norm"] for n in notices})
        record("rederive:off_contract", ok, f"{x['supplier_norm']}: raw total x {factor:g} = "
               f"{tot * factor:.2f}; no council notice supplier normalises to the same string")

    x = first("duplicates")
    if x:
        a, b = resolve(raw, x["row_id_a"]), resolve(raw, x["row_id_b"])
        ok = (a and b and x["row_id_a"] != x["row_id_b"] and a["norm"] == b["norm"]
              and cents(a["amount"]) == cents(b["amount"]) and abs((a["date"] - b["date"]).days) <= 7)
        record("rederive:duplicates", bool(ok), f"{x['row_id_a']} / {x['row_id_b']}")

    x = first("expired_notice_spend")
    if x:
        mine = [n for n in notices if n["norm"] == x["supplier_norm"]]
        ends = [datetime.strptime(n["end"], "%d/%m/%Y").date() if n["end"] else None for n in mine]
        row = [r for r in read_csv("expired_notice_spend.csv") if r["supplier_norm"] == x["supplier_norm"]][0]
        ok = (len(mine) > 0 and all(e is not None and e < window_start for e in ends)
              and any(resolve(raw, i) for i in row["row_ids"].split(";")))
        record("rederive:expired_notice_spend", ok,
               f"{x['supplier_norm']}: {len(mine)} notice block(s), all ended before {window_start}")

    x = first("spend_vs_award")
    if x:
        n = [n for n in notices if n["notice_id"] == x["notice_id"] and n["norm"] == x["supplier_norm"]][0]
        s = datetime.strptime(n["start"], "%d/%m/%Y").date()
        e = datetime.strptime(n["end"], "%d/%m/%Y").date()
        row = [r for r in read_csv("spend_vs_award.csv") if r["supplier_norm"] == x["supplier_norm"]][0]
        tot = sum(resolve(raw, i)["amount"] for i in row["row_ids"].split(";"))
        ratio = tot * factor / (float(n["value"]) / ((e - s).days / 365.25))
        record("rederive:spend_vs_award", round(ratio, 2) == round(x["ratio"], 2),
               f"{x['supplier_norm']} notice {x['notice_id']}: ratio {ratio:.2f} vs {x['ratio']}")

    x = first("high_value_suppliers")
    if x:
        row = [r for r in read_csv("high_value_suppliers.csv") if r["supplier_norm"] == x["supplier_norm"]][0]
        tot = sum(cents(resolve(raw, i)["amount"]) for i in row["row_ids"].split(";"))
        record("rederive:high_value_suppliers", tot == cents(x["window_total"]),
               f"{x['supplier_norm']}: workbook total {tot / 100:.2f} vs {x['window_total']}")

    # 3. every row_id resolves
    bad, n = [], 0
    for p in sorted(glob.glob(os.path.join(FULL, "*.csv"))):
        for r in read_csv(os.path.basename(p)):
            for col in ("row_id", "row_id_a", "row_id_b", "row_ids"):
                if col in r and r[col]:
                    for rid in r[col].split(";"):
                        n += 1
                        if resolve(raw, rid) is None:
                            bad.append(f"{os.path.basename(p)}:{rid}")
    record("row_ids_resolve", not bad, f"{n} row_id references checked; unresolved: {bad[:10]}")

    # 4. no threshold literal in scripts
    hits = []
    # The one permitted literal equal to a T value: a named row count, exempted by name.
    exempt_re = re.compile(r"^CONTRACTS_FINDER_EXPORT_ROW_CAP = \d+\s*(#.*)?$")
    exempted = []
    for T in Ts:
        t = int(T)
        s = str(t)
        grouped = f"{t:,}"
        pats = [r"(?<![\d.])" + re.escape(s) + r"(?:\.0+)?(?![\d])",
                r"(?<![\d.])" + re.escape(grouped).replace(",", "[,_]") + r"(?![\d])"]
        for path in sorted(glob.glob(os.path.join(ROOT, "scripts", "*.py"))):
            with open(path) as f:
                for ln, line in enumerate(f, 1):
                    if os.path.basename(path) == "common.py" and exempt_re.match(line.strip()):
                        exempted.append(f"common.py:{ln}")
                        continue
                    for p in pats:
                        if p and re.search(p, line):
                            hits.append(f"{os.path.basename(path)}:{ln} ({s})")
                    for m in re.finditer(r"(?<![\w.])(\d+(?:\.\d+)?[eE]\d+)\b", line):
                        if float(m.group(1)) == t:
                            hits.append(f"{os.path.basename(path)}:{ln} ({m.group(1)})")
                    # power expressions (base ** exponent)
                    for m in re.finditer(r"(?<![\w.])(\d+)\s*\*\*\s*(\d+)(?![\w.])", line):
                        if int(m.group(1)) ** int(m.group(2)) == t:
                            hits.append(f"{os.path.basename(path)}:{ln} ({m.group(0)})")
                    # k-suffix shorthand: scale the mantissa by three decimal places
                    for m in re.finditer(r"(?<![\w.])(\d+(?:\.\d+)?)[kK]\b", line):
                        if Decimal(m.group(1)).scaleb(3) == t:
                            hits.append(f"{os.path.basename(path)}:{ln} ({m.group(0)})")
    record("no_threshold_literals", not hits, f"T values {sorted(Ts)}; hits: {hits[:10]}; "
           f"exempted by name: {sorted(set(exempted))}")

    # 5. banned words
    bre = re.compile(r"(?<![a-z])(" + "|".join(re.escape(w) for w in BANNED_WORDS) + r")(?![a-z])", re.I)
    with open(AJ) as f:
        found = sorted(set(m.group(1).lower() for m in bre.finditer(f.read())))
    for p in sorted(glob.glob(os.path.join(FULL, "*.csv"))):
        with open(p, newline="") as f:
            hdr = ",".join(next(csv.reader(f), []))
        found += [f"{os.path.basename(p)} header: {m.group(1)}" for m in bre.finditer(hdr)]
    record("banned_words", not found, f"found: {found}")

    status = "PASS" if all(r["result"] == "PASS" for r in results) else "FAIL"
    with open(os.path.join(FULL, "self_check.json"), "w") as f:
        json.dump({"status": status, "checks": results}, f, indent=2, default=str)
    aj["self_check"] = {"status": status, "path": "outputs/analytics-full/self_check.json"}
    with open(AJ, "w") as f:
        json.dump(aj, f, indent=2, default=str)
    for r in results:
        print(f"{r['result']}  {r['check']}: {r['detail'][:160]}")
    print("SELF-CHECK", status)
    sys.exit(0 if status == "PASS" else 1)


if __name__ == "__main__":
    main()
